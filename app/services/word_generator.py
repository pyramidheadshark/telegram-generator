import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def _replace_in_run(run: Run, placeholder: str, value: str, highlight: bool) -> bool:
    """Replace placeholder in a single run, return True if replaced."""
    if placeholder not in run.text:
        return False
    
    new_text = run.text.replace(placeholder, value)
    run.text = new_text
    
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.GREEN
    
    return True


def _replace_in_paragraph(paragraph: Paragraph, replacements: dict[str, str], highlight: bool) -> None:
    """Replace placeholders in paragraph, preserving formatting."""
    full_text = paragraph.text
    if not full_text:
        return
    
    has_replacements = any(p in full_text for p in replacements)
    if not has_replacements:
        return
    
    for placeholder, value in replacements.items():
        if placeholder in full_text:
            for run in paragraph.runs:
                if placeholder in run.text:
                    _replace_in_run(run, placeholder, value, highlight)


def _replace_in_table(table, replacements: dict[str, str], highlight: bool) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, replacements, highlight)


def generate_telegram(
    template_path: str,
    data: dict,
    output_path: str,
    highlight: bool = True,
) -> str:
    template = Path(template_path)
    if not template.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")

    doc = Document(template_path)

    replacements = {
        "{{адрес}}": data.get("address", ""),
        "{{организация}}": data.get("organization", ""),
        "{{должность_дат}}": data.get("position_dative", ""),
        "{{фио_дат}}": data.get("fullname_dative", ""),
        "{{имя_отчество}}": data.get("name_for_greeting", ""),
        "{{уважаемый}}": data.get("greeting", ""),
    }

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, replacements, highlight)

    for table in doc.tables:
        _replace_in_table(table, replacements, highlight)

    for section in doc.sections:
        for header in [section.header, section.footer]:
            for paragraph in header.paragraphs:
                _replace_in_paragraph(paragraph, replacements, highlight)
            for table in header.tables:
                _replace_in_table(table, replacements, highlight)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    return str(output.absolute())


def validate_template(template_path: str) -> list[str]:
    template = Path(template_path)
    if not template.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")

    doc = Document(template_path)

    all_text = ""
    for paragraph in doc.paragraphs:
        all_text += paragraph.text + " "

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    all_text += paragraph.text + " "

    for section in doc.sections:
        for header in [section.header, section.footer]:
            for paragraph in header.paragraphs:
                all_text += paragraph.text + " "

    pattern = r"\{\{[^}]+\}\}"
    found_placeholders = re.findall(pattern, all_text)

    required_placeholders = {
        "{{адрес}}",
        "{{организация}}",
        "{{должность_дат}}",
        "{{фио_дат}}",
        "{{имя_отчество}}",
        "{{уважаемый}}",
    }

    found_set = set(found_placeholders)
    missing = required_placeholders - found_set
    unknown = found_set - required_placeholders

    errors = []
    if missing:
        errors.append(f"Missing placeholders: {', '.join(sorted(missing))}")
    if unknown:
        errors.append(f"Unknown placeholders: {', '.join(sorted(unknown))}")

    return errors