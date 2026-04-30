from pathlib import Path

from openpyxl import load_workbook


def parse_excel(file_path: str) -> list[dict]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = workbook.active

    if sheet is None:
        workbook.close()
        raise ValueError("Excel file has no active sheet")

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        workbook.close()
        raise ValueError("Excel file is empty or has no header row")

    header_map = {}
    for idx, cell in enumerate(header_row):
        if cell:
            header_map[idx] = str(cell).strip()

    column_mapping = {
        "Организация": "organization",
        "ФИО": "fullname",
        "Должность": "position",
        "Адрес": "address",
        "Ответственный": "responsible",
    }

    col_indices = {}
    for idx, header in header_map.items():
        if header in column_mapping:
            col_indices[column_mapping[header]] = idx

    required_columns = ["organization", "fullname", "position", "address"]
    missing = [col for col in required_columns if col not in col_indices]
    if missing:
        workbook.close()
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    results = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or all(cell is None for cell in row):
            continue

        record = {}
        for col_name, col_idx in col_indices.items():
            if col_idx < len(row):
                value = row[col_idx]
                record[col_name] = str(value).strip() if value is not None else ""
            else:
                record[col_name] = ""

        if not record.get("fullname") or not record.get("organization"):
            continue

        results.append(record)

    workbook.close()
    return results
